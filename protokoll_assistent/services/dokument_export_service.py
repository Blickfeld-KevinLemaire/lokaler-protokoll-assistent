"""Export von Transkript und Protokoll in gaengige Formate.

Aus den Ergebnisdateien der Verarbeitung (Transkript-JSON, Protokoll-JSON)
wird erst ein einfaches, formatfreies **Dokument** gebaut (Titel, Kopfzeilen,
Abschnitte) und dann in das gewuenschte Format geschrieben:

    Word (.docx), PDF, Markdown, Text, HTML, OpenDocument (.odt),
    dazu fuer das Transkript Untertitel (SRT, VTT) und JSON.

Zwei Wege nutzen das:

* **Automatisch**: Nach jeder Verarbeitung legt ``automatisches_word`` eine
  zusammengefasste Word-Datei neben die uebrigen Ausgaben (Protokoll und
  Transkript in einem Dokument, ohne Protokoll nur das Transkript).
* **Auf Wunsch**: Der Exportdialog laesst den Anwender Inhalt, Formate und
  Zielordner waehlen (``exportiere``).

Der Dienst kennt kein Qt (CLAUDE.md). PDF und OpenDocument erzeugt Qt aus dem
HTML-Text; diese Schreiber werden von der Oberflaeche als Funktionen
uebergeben (``schreiber``), genau wie die ML-Aufrufe der anderen Dienste.
Damit bleibt der Dienst ohne Fenster testbar und braucht keine zusaetzliche
Bibliothek fuer PDF.
"""

from __future__ import annotations

import html
import json
import re
import tempfile
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from protokoll_assistent.services import export_service
from protokoll_assistent.services.manifest_service import _ersetzen_mit_wiederholung

INHALT_TRANSKRIPT = "transkript"
INHALT_PROTOKOLL = "protokoll"
INHALT_BEIDES = "beides"

# Schluessel in der Protokoll-JSON, unter dem der in der Vorschau korrigierte Text steht.
SCHLUESSEL_KORRIGIERT = "korrigierter_text"

# Abschnittsarten eines Dokuments
UEBERSCHRIFT = "ueberschrift"
ABSATZ = "absatz"
LISTE = "liste"
SEITENUMBRUCH = "seitenumbruch"


class ExportFehler(RuntimeError):
    """Export nicht moeglich, mit lesbarer Meldung."""


@dataclass(frozen=True)
class Format:
    id: str
    name: str
    endung: str
    inhalte: tuple[str, ...] = (INHALT_TRANSKRIPT, INHALT_PROTOKOLL, INHALT_BEIDES)


FORMATE: tuple[Format, ...] = (
    Format("docx", "Word (.docx)", ".docx"),
    Format("pdf", "PDF (.pdf)", ".pdf"),
    Format("md", "Markdown (.md)", ".md"),
    Format("txt", "Text (.txt)", ".txt"),
    Format("html", "HTML (.html)", ".html"),
    Format("odt", "OpenDocument (.odt)", ".odt"),
    Format("srt", "Untertitel SRT (.srt)", ".srt", (INHALT_TRANSKRIPT,)),
    Format("vtt", "Untertitel VTT (.vtt)", ".vtt", (INHALT_TRANSKRIPT,)),
    Format("json", "JSON (.json)", ".json", (INHALT_TRANSKRIPT, INHALT_PROTOKOLL)),
)

# Formate, die Qt aus dem HTML-Text schreibt (siehe 'gui/dokument_qt.py').
QT_FORMATE = ("pdf", "odt")

Schreiber = Callable[[str, Path], None]  # (HTML-Text, Zielpfad)


def finde_format(format_id: str) -> Format | None:
    return next((f for f in FORMATE if f.id == format_id), None)


def formate_fuer(inhalt: str) -> list[Format]:
    return [f for f in FORMATE if inhalt in f.inhalte]


# --------------------------------------------------------------------------
# Dokumentmodell
# --------------------------------------------------------------------------
@dataclass(frozen=True)
class Abschnitt:
    art: str
    text: str = ""
    stufe: int = 1  # nur Ueberschriften: 1 = Titel, 2 = Abschnitt
    punkte: tuple[str, ...] = ()  # nur Listen
    betont: str = ""  # nur Absaetze: fett gesetzter Anfang ("[00:01] Anna:")


@dataclass
class Dokument:
    titel: str
    kopfzeilen: list[str] = field(default_factory=list)
    abschnitte: list[Abschnitt] = field(default_factory=list)


# --------------------------------------------------------------------------
# Dateien lesen
# --------------------------------------------------------------------------
def _lade_json(pfad: Path, was: str) -> dict[str, Any]:
    try:
        daten = json.loads(pfad.read_text(encoding="utf-8"))
    except (OSError, ValueError) as fehler:
        raise ExportFehler(f"{was} konnte nicht gelesen werden ({pfad.name}): {fehler}") from fehler
    if not isinstance(daten, dict):
        raise ExportFehler(f"{was} ({pfad.name}) hat nicht das erwartete Format.")
    return daten


def lade_transkript(pfad: Path) -> dict[str, Any]:
    daten = _lade_json(pfad, "Das Transkript")
    if not isinstance(daten.get("segmente"), list):
        raise ExportFehler(f"Das Transkript ({pfad.name}) enthaelt keine Segmente.")
    return daten


def lade_protokoll(pfad: Path) -> dict[str, Any]:
    return _lade_json(pfad, "Das Protokoll")


# --------------------------------------------------------------------------
# Dokumente bauen
# --------------------------------------------------------------------------
def transkript_dokument(daten: dict[str, Any]) -> Dokument:
    quelle = str(daten.get("quelldatei") or "Aufnahme")
    kopf = [f"Quelldatei: {quelle}"]
    if daten.get("erstellt"):
        kopf.append(f"Erstellt: {daten['erstellt']}")
    if daten.get("modell"):
        kopf.append(f"Transkriptionsmodell: {daten['modell']}")
    if daten.get("sprache"):
        kopf.append(f"Sprache: {daten['sprache']}")
    mit_sprechern = bool(daten.get("sprechertrennung_aktiv", True))
    if mit_sprechern:
        kopf.append(f"Erkannte Sprecher: {daten.get('anzahl_sprecher', 0)}")

    dokument = Dokument(f"Transkript: {quelle}", kopf)
    for segment in daten["segmente"]:
        if not isinstance(segment, dict):
            continue
        zeit = f"[{segment.get('start', '')}]"
        text = str(segment.get("text", ""))
        if mit_sprechern:
            dokument.abschnitte.append(Abschnitt(ABSATZ, f" {text}", betont=f"{zeit} {segment.get('sprecher', '')}:"))
        else:
            dokument.abschnitte.append(Abschnitt(ABSATZ, f" {text}", betont=zeit))
    return dokument


def protokoll_dokument(protokoll: dict[str, Any], mit_korrektur: bool = True) -> Dokument:
    """Dieselben Abschnitte wie ``export_service.render_protocol_markdown``.

    Hat der Anwender das Protokoll in der Vorschau korrigiert, steht sein Text
    unter ``SCHLUESSEL_KORRIGIERT`` in der Protokoll-JSON (die Auswertung des
    Modells bleibt daneben unveraendert erhalten). Jeder Export geht dann von
    diesem Text aus -- es wird exportiert, was in der Vorschau steht.
    ``mit_korrektur=False`` liefert die urspruengliche Auswertung."""
    korrigiert = protokoll.get(SCHLUESSEL_KORRIGIERT)
    if mit_korrektur and isinstance(korrigiert, str) and korrigiert.strip():
        return dokument_aus_markdown(korrigiert)
    dokument = Dokument(str(protokoll.get("titel") or "Protokoll"))
    zusammenfassung = str(protokoll.get("kurzzusammenfassung") or "").strip()
    if zusammenfassung:
        dokument.abschnitte.append(Abschnitt(ABSATZ, zusammenfassung))

    def abschnitt(titel: str, eintraege: Any, formatierer: Callable[[Any], str]) -> None:
        if not isinstance(eintraege, list) or not eintraege:
            return
        dokument.abschnitte.append(Abschnitt(UEBERSCHRIFT, titel, stufe=2))
        dokument.abschnitte.append(Abschnitt(LISTE, punkte=tuple(formatierer(e) for e in eintraege)))

    def unklar(wert: Any) -> str:
        return str(wert) if wert else "unklar"

    abschnitt(
        "Themen",
        protokoll.get("themen"),
        lambda t: f"{t.get('thema', '')} ({t.get('zeitraum', '')}): " + "; ".join(t.get("kernaussagen", [])),
    )
    abschnitt(
        "Entscheidungen",
        protokoll.get("entscheidungen"),
        lambda e: f"{e.get('entscheidung', '')} (Sprecher: {unklar(e.get('sprecher'))}, "
        f"Zeitpunkt: {unklar(e.get('zeitpunkt'))}, Quelle: {e.get('quelle', '')})",
    )
    abschnitt(
        "Aufgaben",
        protokoll.get("aufgaben"),
        lambda a: f"{a.get('aufgabe', '')} (Verantwortlich: {unklar(a.get('verantwortlich'))}, "
        f"Frist: {unklar(a.get('frist'))}, Quelle: {a.get('quelle', '')})",
    )
    for titel, schluessel in (
        ("Termine", "termine"),
        ("Offene Fragen", "offene_fragen"),
        ("Wichtige Fakten", "wichtige_fakten"),
        ("Unsichere Transkriptstellen", "unsichere_transkriptstellen"),
        ("Quellenhinweise", "quellenhinweise"),
    ):
        abschnitt(titel, protokoll.get(schluessel), export_service.listeneintrag_als_text)
    return dokument


def dokument_als_markdown(dokument: Dokument) -> str:
    """Ein Dokument als schlichter, bearbeitbarer Text: ``# Titel``, Absaetze,
    ``## Abschnitt`` und ``- Punkt``. Das ist genau das, was die Vorschau zeigt
    und ``dokument_aus_markdown`` wieder einliest."""
    zeilen = [f"# {dokument.titel}", ""]
    for abschnitt in dokument.abschnitte:
        if abschnitt.art == UEBERSCHRIFT:
            zeilen += [f"{'#' * max(abschnitt.stufe, 1)} {abschnitt.text}", ""]
        elif abschnitt.art == LISTE:
            zeilen += [f"- {punkt}" for punkt in abschnitt.punkte] + [""]
        elif abschnitt.art == ABSATZ:
            zeilen += [(f"{abschnitt.betont}{abschnitt.text}").strip(), ""]
    return "\n".join(zeilen).rstrip() + "\n"


_FETT = re.compile(r"\*\*(.+?)\*\*")
_UEBERSCHRIFT = re.compile(r"^(#{1,6})\s+(.*\S)\s*$")
_LISTENPUNKT = re.compile(r"^[-*•]\s+(.*\S)\s*$")


def dokument_aus_markdown(text: str) -> Dokument:
    """Liest den Text der Vorschau zurueck. Erlaubt ist, was ``dokument_als_markdown``
    schreibt: die erste ``# ``-Zeile ist der Titel, weitere Ueberschriften sind
    Abschnitte, ``- `` ist ein Listenpunkt, jede andere Zeile ein eigener Absatz.
    Hervorhebungen mit ``**`` fallen weg (das Dokumentmodell kennt sie nur im
    Transkript). Leerzeilen trennen nur und zaehlen nicht."""
    dokument = Dokument("Protokoll")
    titel_gesetzt = False
    punkte: list[str] = []

    def liste_abschliessen() -> None:
        if punkte:
            dokument.abschnitte.append(Abschnitt(LISTE, punkte=tuple(punkte)))
            punkte.clear()

    for zeile in text.splitlines():
        zeile = _FETT.sub(r"\1", zeile.strip())
        if not zeile:
            liste_abschliessen()
            continue
        if (treffer := _UEBERSCHRIFT.match(zeile)) is not None:
            liste_abschliessen()
            if treffer.group(1) == "#" and not titel_gesetzt:
                dokument.titel = treffer.group(2)
                titel_gesetzt = True
            else:
                dokument.abschnitte.append(Abschnitt(UEBERSCHRIFT, treffer.group(2), stufe=2))
        elif (punkt := _LISTENPUNKT.match(zeile)) is not None:
            punkte.append(punkt.group(1))
        else:
            liste_abschliessen()
            dokument.abschnitte.append(Abschnitt(ABSATZ, zeile))
    liste_abschliessen()
    return dokument


def kombiniertes_dokument(protokoll: Dokument, transkript: Dokument) -> Dokument:
    """Protokoll vorn, das Transkript als Anhang auf einer neuen Seite."""
    dokument = Dokument(protokoll.titel, [*protokoll.kopfzeilen, *transkript.kopfzeilen])
    dokument.abschnitte.extend(protokoll.abschnitte)
    dokument.abschnitte.append(Abschnitt(SEITENUMBRUCH))
    dokument.abschnitte.append(Abschnitt(UEBERSCHRIFT, "Transkript", stufe=2))
    dokument.abschnitte.extend(transkript.abschnitte)
    return dokument


# --------------------------------------------------------------------------
# Renderer
# --------------------------------------------------------------------------
def als_markdown(dokument: Dokument) -> str:
    zeilen = [f"# {dokument.titel}", ""]
    zeilen += [f"{kopf}  " for kopf in dokument.kopfzeilen]
    if dokument.kopfzeilen:
        zeilen.append("")
    for a in dokument.abschnitte:
        if a.art == UEBERSCHRIFT:
            zeilen += [f"{'#' * a.stufe} {a.text}", ""]
        elif a.art == ABSATZ:
            zeilen += [f"**{a.betont}**{a.text}" if a.betont else a.text, ""]
        elif a.art == LISTE:
            zeilen += [f"- {punkt}" for punkt in a.punkte]
            zeilen.append("")
        elif a.art == SEITENUMBRUCH:
            zeilen += ["---", ""]
    return "\n".join(zeilen).rstrip("\n") + "\n"


def als_text(dokument: Dokument) -> str:
    zeilen = [dokument.titel.upper(), "=" * len(dokument.titel), ""]
    zeilen += dokument.kopfzeilen
    if dokument.kopfzeilen:
        zeilen.append("")
    for a in dokument.abschnitte:
        if a.art == UEBERSCHRIFT:
            zeilen += [a.text, "-" * len(a.text)]
        elif a.art == ABSATZ:
            zeilen.append(f"{a.betont}{a.text}")
        elif a.art == LISTE:
            zeilen += [f"  - {punkt}" for punkt in a.punkte]
            zeilen.append("")
        elif a.art == SEITENUMBRUCH:
            zeilen += ["", "*" * 40, ""]
    return "\n".join(zeilen).rstrip("\n") + "\n"


def als_html(dokument: Dokument) -> str:
    """Eigenstaendige HTML-Seite. Auch die Quelle fuer PDF und OpenDocument."""
    e = html.escape
    teile = [
        "<!DOCTYPE html>",
        '<html lang="de"><head><meta charset="utf-8">',
        f"<title>{e(dokument.titel)}</title>",
        (
            "<style>body{font-family:'Segoe UI',Arial,sans-serif;line-height:1.45;margin:2em;}"
            "h1{font-size:20pt;} h2{font-size:14pt;margin-top:1.4em;} .kopf{color:#555;margin:0;}"
            "li{margin:0.2em 0;}</style></head><body>"
        ),
        f"<h1>{e(dokument.titel)}</h1>",
    ]
    teile += [f'<p class="kopf">{e(kopf)}</p>' for kopf in dokument.kopfzeilen]
    for a in dokument.abschnitte:
        if a.art == UEBERSCHRIFT:
            teile.append(f"<h{a.stufe}>{e(a.text)}</h{a.stufe}>")
        elif a.art == ABSATZ:
            vorn = f"<strong>{e(a.betont)}</strong>" if a.betont else ""
            teile.append(f"<p>{vorn}{e(a.text)}</p>")
        elif a.art == LISTE:
            teile.append("<ul>" + "".join(f"<li>{e(punkt)}</li>" for punkt in a.punkte) + "</ul>")
        elif a.art == SEITENUMBRUCH:
            teile.append('<p style="page-break-after: always"></p>')
    teile.append("</body></html>")
    return "\n".join(teile) + "\n"


def schreibe_docx(dokument: Dokument, pfad: Path) -> None:
    try:
        import docx  # type: ignore
    except ImportError as fehler:
        raise ExportFehler("Fuer Word-Dateien wird das Paket python-docx benoetigt.") from fehler

    word = docx.Document()
    word.core_properties.title = dokument.titel
    word.add_heading(dokument.titel, level=1)
    for kopf in dokument.kopfzeilen:
        word.add_paragraph(kopf)
    for a in dokument.abschnitte:
        if a.art == UEBERSCHRIFT:
            word.add_heading(a.text, level=a.stufe)
        elif a.art == ABSATZ:
            absatz = word.add_paragraph()
            if a.betont:
                absatz.add_run(a.betont).bold = True
            absatz.add_run(a.text)
        elif a.art == LISTE:
            for punkt in a.punkte:
                word.add_paragraph(punkt, style="List Bullet")
        elif a.art == SEITENUMBRUCH:
            word.add_page_break()
    word.save(str(pfad))


# --------------------------------------------------------------------------
# Schreiben
# --------------------------------------------------------------------------
def eindeutiger_pfad(ordner: Path, basis: str, endung: str) -> Path:
    """Ueberschreibt nie eine vorhandene Datei: ``name.ext``, ``name_2.ext`` ..."""
    kandidat = ordner / f"{basis}{endung}"
    zaehler = 2
    while kandidat.exists():
        kandidat = ordner / f"{basis}_{zaehler}{endung}"
        zaehler += 1
    return kandidat


def _atomar_schreiben(ziel: Path, schreiben: Callable[[Path], None]) -> None:
    """Erst unter Arbeitsnamen schreiben, dann ersetzen -- mit Wiederholung,
    weil Windows (Virenscanner, Suchindex, ein geoeffnetes Word) das Ersetzen
    kurz mit 'Zugriff verweigert' ablehnen kann (CLAUDE.md, Regel 1)."""
    ziel.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=ziel.parent, suffix=".tmp", delete=False) as handle:
        temp = Path(handle.name)
    try:
        schreiben(temp)
        _ersetzen_mit_wiederholung(temp, ziel)
    finally:
        temp.unlink(missing_ok=True)


def _basisname(transkript: dict[str, Any] | None, protokoll_pfad: Path | None, transkript_pfad: Path | None) -> str:
    if transkript is not None and transkript.get("quelldatei_stamm"):
        return str(transkript["quelldatei_stamm"])
    for pfad in (transkript_pfad, protokoll_pfad):
        if pfad is not None:
            return pfad.stem
    return "export"


def dokument_fuer(
    inhalt: str,
    transkript_json: Path | None,
    protokoll_json: Path | None,
) -> tuple[Dokument, dict[str, Any] | None, dict[str, Any] | None]:
    """Baut das Dokument fuer den gewaehlten Inhalt und liefert dazu die gelesenen Rohdaten."""
    transkript = lade_transkript(transkript_json) if transkript_json is not None else None
    protokoll = lade_protokoll(protokoll_json) if protokoll_json is not None else None
    if inhalt == INHALT_TRANSKRIPT:
        if transkript is None:
            raise ExportFehler("Es ist kein Transkript ausgewaehlt.")
        return transkript_dokument(transkript), transkript, protokoll
    if inhalt == INHALT_PROTOKOLL:
        if protokoll is None:
            raise ExportFehler("Es liegt noch kein Protokoll vor.")
        return protokoll_dokument(protokoll), transkript, protokoll
    if inhalt == INHALT_BEIDES:
        if transkript is None or protokoll is None:
            raise ExportFehler("Fuer das gemeinsame Dokument werden Transkript und Protokoll gebraucht.")
        return kombiniertes_dokument(protokoll_dokument(protokoll), transkript_dokument(transkript)), transkript, protokoll
    raise ExportFehler(f"Unbekannter Inhalt: {inhalt}")


def exportiere(
    inhalt: str,
    format_ids: list[str],
    zielordner: Path,
    *,
    transkript_json: Path | None = None,
    protokoll_json: Path | None = None,
    schreiber: dict[str, Schreiber] | None = None,
) -> list[Path]:
    """Schreibt das Dokument in alle gewaehlten Formate nach ``zielordner``.

    Gibt die geschriebenen Dateien zurueck. Vorhandene Dateien werden nie
    ueberschrieben (es wird hochgezaehlt). Schlaegt ein Format fehl, wird mit
    den uebrigen weitergemacht und am Ende alles Gescheiterte gemeldet --
    sonst ginge wegen eines einzigen Formats der ganze Export verloren."""
    if not format_ids:
        raise ExportFehler("Bitte mindestens ein Format waehlen.")
    dokument, transkript, protokoll = dokument_fuer(inhalt, transkript_json, protokoll_json)
    schreiber = schreiber or {}
    zielordner.mkdir(parents=True, exist_ok=True)
    basis = f"{_basisname(transkript, protokoll_json, transkript_json)}_{inhalt if inhalt != INHALT_BEIDES else 'gesamt'}"

    geschrieben: list[Path] = []
    fehler: list[str] = []
    for format_id in format_ids:
        format_ = finde_format(format_id)
        if format_ is None or inhalt not in format_.inhalte:
            fehler.append(f"{format_id}: fuer diesen Inhalt nicht moeglich")
            continue
        ziel = eindeutiger_pfad(zielordner, basis, format_.endung)
        try:
            _atomar_schreiben(
                ziel, _schreib_funktion(format_id, inhalt, dokument, transkript, protokoll, schreiber)
            )
        except (ExportFehler, OSError, RuntimeError) as problem:
            fehler.append(f"{format_.name}: {problem}")
            continue
        geschrieben.append(ziel)
    if fehler and not geschrieben:
        raise ExportFehler("Der Export ist fehlgeschlagen:\n" + "\n".join(fehler))
    if fehler:
        raise ExportTeilweise(geschrieben, fehler)
    return geschrieben


def _schreib_funktion(
    format_id: str,
    inhalt: str,
    dokument: Dokument,
    transkript: dict[str, Any] | None,
    protokoll: dict[str, Any] | None,
    schreiber: dict[str, Schreiber],
) -> Callable[[Path], None]:
    def schreiben(temp: Path) -> None:
        _schreibe_format(format_id, inhalt, dokument, temp, transkript, protokoll, schreiber)

    return schreiben


class ExportTeilweise(ExportFehler):
    """Ein Teil der Formate wurde geschrieben, andere nicht."""

    def __init__(self, geschrieben: list[Path], fehler: list[str]):
        super().__init__("Nicht alle Formate konnten geschrieben werden:\n" + "\n".join(fehler))
        self.geschrieben = geschrieben
        self.fehler = fehler


def _schreibe_format(
    format_id: str,
    inhalt: str,
    dokument: Dokument,
    ziel: Path,
    transkript: dict[str, Any] | None,
    protokoll: dict[str, Any] | None,
    schreiber: dict[str, Schreiber],
) -> None:
    if format_id == "docx":
        schreibe_docx(dokument, ziel)
    elif format_id == "md":
        ziel.write_text(als_markdown(dokument), encoding="utf-8")
    elif format_id == "txt":
        ziel.write_text(als_text(dokument), encoding="utf-8")
    elif format_id == "html":
        ziel.write_text(als_html(dokument), encoding="utf-8")
    elif format_id in QT_FORMATE:
        funktion = schreiber.get(format_id)
        if funktion is None:
            raise ExportFehler("Dieses Format steht in dieser Umgebung nicht zur Verfuegung.")
        funktion(als_html(dokument), ziel)
    elif format_id in ("srt", "vtt"):
        if transkript is None:
            raise ExportFehler("Untertitel brauchen ein Transkript.")
        segmente = [
            {
                "start": segment["start_sekunden"],
                "end": segment["ende_sekunden"],
                "sprecher": segment.get("sprecher", ""),
                "text": segment["text"],
            }
            for segment in transkript["segmente"]
        ]
        mit_sprechern = bool(transkript.get("sprechertrennung_aktiv", True))
        baue = export_service.build_srt_content if format_id == "srt" else export_service.build_vtt_content
        ziel.write_text(baue(segmente, mit_sprechern), encoding="utf-8")
    elif format_id == "json":
        quelle = transkript if inhalt == INHALT_TRANSKRIPT else protokoll
        ziel.write_text(json.dumps(quelle, ensure_ascii=False, indent=2), encoding="utf-8")
    else:  # pragma: no cover - 'finde_format' kennt nur die Formate oben
        raise ExportFehler(f"Unbekanntes Format: {format_id}")


# --------------------------------------------------------------------------
# Automatische Word-Datei
# --------------------------------------------------------------------------
def automatisches_word(transkript_json: Path, protokoll_json: Path | None = None) -> Path | None:
    """Die zusammengefasste Word-Datei, die nach jeder Verarbeitung von selbst entsteht.

    Mit Protokoll: ein Dokument aus Protokoll und Transkript, daneben das
    Protokoll (gleicher Name wie dessen JSON, Endung ``.docx``). Ohne Protokoll:
    das Transkript neben seiner JSON. Eine vorhandene Datei wird ersetzt --
    sie gehoert zum selben Lauf und wird bei umbenannten Sprechern neu erzeugt.

    Gibt ``None`` zurueck, wenn ``python-docx`` fehlt; andere Fehler werden als
    ``ExportFehler`` gemeldet (der Aufrufer faengt sie ab: die Verarbeitung
    selbst darf daran nie scheitern)."""
    try:
        import docx  # type: ignore  # noqa: F401
    except ImportError:
        return None
    if protokoll_json is not None and protokoll_json.is_file():
        dokument, _, _ = dokument_fuer(INHALT_BEIDES, transkript_json, protokoll_json)
        ziel = protokoll_json.with_suffix(".docx")
    else:
        dokument, _, _ = dokument_fuer(INHALT_TRANSKRIPT, transkript_json, None)
        ziel = transkript_json.with_suffix(".docx")
    def schreiben(temp: Path) -> None:
        schreibe_docx(dokument, temp)

    try:
        _atomar_schreiben(ziel, schreiben)
    except OSError as fehler:
        raise ExportFehler(f"Die Word-Datei konnte nicht geschrieben werden: {fehler}") from fehler
    return ziel
